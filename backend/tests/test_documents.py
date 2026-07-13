from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import httpx2 as httpx
from docx import Document
from fastapi.testclient import TestClient
from pydantic import SecretStr
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.config import Settings
from app.main import create_app

INTERNAL_TOKEN = "test-internal-token-that-is-long-enough"
PDF_MEDIA_TYPE = "application/pdf"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_FILE_BYTES = 10 * 1024 * 1024


def make_pdf(
    page_texts: list[str],
    *,
    password: str | None = None,
    attachment: bytes | None = None,
) -> bytes:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    for text in page_texts:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        escaped_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content = DecodedStreamObject()
        content.set_data(f"BT /F1 12 Tf 72 720 Td ({escaped_text}) Tj ET".encode())
        page[NameObject("/Contents")] = content

    if password is not None:
        writer.encrypt(password)
    if attachment is not None:
        writer.add_attachment("padding.bin", attachment)

    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx(*paragraphs: str, table_cells: tuple[str, str] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_cells is not None:
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = table_cells[0]
        table.cell(0, 1).text = table_cells[1]
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_zip(entries: dict[str, bytes]) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", compression=ZIP_DEFLATED) as archive:
        for name, value in entries.items():
            archive.writestr(name, value)
    return buffer.getvalue()


def read_zip_entries(content: bytes) -> dict[str, bytes]:
    with ZipFile(BytesIO(content)) as archive:
        return {info.filename: archive.read(info) for info in archive.infolist()}


def make_pdf_at_size(page_texts: list[str], size: int) -> bytes:
    payload_size = size - len(make_pdf(page_texts)) - 1_024
    for _ in range(5):
        content = make_pdf(page_texts, attachment=b"x" * payload_size)
        difference = size - len(content)
        if difference == 0:
            return content
        payload_size += difference
    raise AssertionError("could not construct an exact-size PDF fixture")


def post_document(
    content: bytes,
    *,
    filename: str,
    media_type: str,
    kind: str = "resume",
    include_token: bool = True,
) -> httpx.Response:
    headers = {"X-GradPath-Internal-Token": INTERNAL_TOKEN} if include_token else {}
    settings = Settings(internal_api_token=SecretStr(INTERNAL_TOKEN))
    with TestClient(create_app(settings=settings)) as client:
        return client.post(
            "/api/v1/documents/extract",
            data={"kind": kind},
            files={"file": (filename, content, media_type)},
            headers=headers,
        )


def test_extracts_native_pdf_text() -> None:
    content = make_pdf(["Built a FastAPI service for campus placements."])

    response = post_document(
        content,
        filename="resume.pdf",
        media_type=PDF_MEDIA_TYPE,
    )

    assert response.status_code == 200
    result = response.json()
    assert set(result) == {
        "documentKind",
        "documentFormat",
        "text",
        "pageCount",
        "characterCount",
        "extractionMethod",
        "layoutPreserved",
        "extractionVersion",
    }
    assert result["documentKind"] == "resume"
    assert result["documentFormat"] == "pdf"
    assert result["pageCount"] == 1
    assert "Built a FastAPI service for campus placements." in result["text"]
    assert result["characterCount"] == len(result["text"])
    assert result["extractionMethod"] == "native_text"
    assert result["layoutPreserved"] is False
    assert result["extractionVersion"] == "native-text-v1"
    assert response.headers["Cache-Control"] == "no-store"


def test_extracts_docx_paragraphs_and_tables_in_document_order() -> None:
    content = make_docx(
        "Software intern role",
        table_cells=("Required skill", "PostgreSQL"),
    )

    response = post_document(
        content,
        filename="role.docx",
        media_type=DOCX_MEDIA_TYPE,
        kind="jd",
    )

    assert response.status_code == 200
    result = response.json()
    assert result["documentKind"] == "jd"
    assert result["documentFormat"] == "docx"
    assert result["pageCount"] is None
    assert result["text"].index("Software intern role") < result["text"].index("Required skill")
    assert "Required skill\tPostgreSQL" in result["text"]


def test_document_extraction_requires_internal_authentication() -> None:
    response = post_document(
        make_pdf(["Private resume text must not be parsed without authentication."]),
        filename="resume.pdf",
        media_type=PDF_MEDIA_TYPE,
        include_token=False,
    )

    assert response.status_code == 401


def test_document_extraction_rejects_unknown_kind() -> None:
    response = post_document(
        make_pdf(["A valid PDF body with enough native text for extraction."]),
        filename="resume.pdf",
        media_type=PDF_MEDIA_TYPE,
        kind="transcript",
    )

    assert response.status_code == 422


def test_document_extraction_requires_a_file() -> None:
    settings = Settings(internal_api_token=SecretStr(INTERNAL_TOKEN))
    with TestClient(create_app(settings=settings)) as client:
        response = client.post(
            "/api/v1/documents/extract",
            data={"kind": "resume"},
            headers={"X-GradPath-Internal-Token": INTERNAL_TOKEN},
        )

    assert response.status_code == 422
    assert response.json() == {"detail": "Exactly one document file is required"}


def test_document_extraction_rejects_an_empty_file() -> None:
    response = post_document(
        b"",
        filename="empty.pdf",
        media_type=PDF_MEDIA_TYPE,
    )

    assert response.status_code == 422
    assert response.json() == {"detail": "Uploaded file is empty"}


def test_document_extraction_rejects_unsupported_type() -> None:
    response = post_document(
        b"This is plain text, not an accepted document container.",
        filename="resume.txt",
        media_type="text/plain",
    )

    assert response.status_code == 415
    assert response.json() == {"detail": "Unsupported document type"}
    assert response.headers["Cache-Control"] == "no-store"


def test_document_extraction_rejects_mismatched_magic_bytes() -> None:
    response = post_document(
        make_docx("This package is DOCX content."),
        filename="resume.pdf",
        media_type=PDF_MEDIA_TYPE,
    )

    assert response.status_code == 415
    assert response.json() == {"detail": "Document type does not match its content"}


def test_document_extraction_rejects_files_over_ten_mibibytes() -> None:
    content = b"%PDF-1.7\n".ljust(MAX_FILE_BYTES + 1, b"x")

    response = post_document(
        content,
        filename="oversized.pdf",
        media_type=PDF_MEDIA_TYPE,
    )

    assert response.status_code == 413
    assert response.json() == {"detail": "Uploaded file exceeds the 10 MiB limit"}


def test_document_extraction_accepts_a_file_at_exact_size_limit() -> None:
    content = make_pdf_at_size(
        ["A native-text PDF exactly at the accepted file-size boundary."],
        MAX_FILE_BYTES,
    )

    response = post_document(
        content,
        filename="boundary.pdf",
        media_type=PDF_MEDIA_TYPE,
    )

    assert response.status_code == 200


def test_document_extraction_enforces_resume_page_limit() -> None:
    content = make_pdf([f"Resume page {index} has native text." for index in range(11)])

    response = post_document(
        content,
        filename="long-resume.pdf",
        media_type=PDF_MEDIA_TYPE,
    )

    assert response.status_code == 413
    assert response.json() == {"detail": "Document exceeds the resume page limit of 10"}


def test_document_extraction_allows_eleven_page_jd() -> None:
    content = make_pdf([f"Job description page {index} has native text." for index in range(11)])

    response = post_document(
        content,
        filename="jd.pdf",
        media_type=PDF_MEDIA_TYPE,
        kind="jd",
    )

    assert response.status_code == 200
    assert response.json()["pageCount"] == 11


def test_document_extraction_enforces_jd_page_limit() -> None:
    content = make_pdf([f"Job description page {index} has native text." for index in range(31)])

    response = post_document(
        content,
        filename="long-jd.pdf",
        media_type=PDF_MEDIA_TYPE,
        kind="jd",
    )

    assert response.status_code == 413
    assert response.json() == {"detail": "Document exceeds the JD page limit of 30"}


def test_document_extraction_rejects_encrypted_pdf() -> None:
    content = make_pdf(["Confidential resume text."], password="secret")

    response = post_document(
        content,
        filename="encrypted.pdf",
        media_type=PDF_MEDIA_TYPE,
    )

    assert response.status_code == 422
    assert response.json() == {"detail": "Encrypted documents are not supported"}


def test_document_extraction_rejects_pdf_without_native_text() -> None:
    content = make_pdf([""])

    response = post_document(
        content,
        filename="scanned.pdf",
        media_type=PDF_MEDIA_TYPE,
    )

    assert response.status_code == 422
    assert response.json() == {
        "detail": "Document has no machine-readable text; OCR is not enabled"
    }


def test_document_extraction_rejects_corrupt_pdf_with_sanitized_error() -> None:
    response = post_document(
        b"%PDF-1.7\nnot a valid PDF body",
        filename="corrupt.pdf",
        media_type=PDF_MEDIA_TYPE,
    )

    assert response.status_code == 422
    assert response.json() == {"detail": "Document could not be parsed"}


def test_document_extraction_rejects_corrupt_docx_with_sanitized_error() -> None:
    response = post_document(
        b"PK\x03\x04truncated package",
        filename="corrupt.docx",
        media_type=DOCX_MEDIA_TYPE,
    )

    assert response.status_code == 422
    assert response.json() == {"detail": "Document could not be parsed"}


def test_document_extraction_rejects_generic_zip_disguised_as_docx() -> None:
    response = post_document(
        make_zip({"notes.txt": b"This is a generic archive, not a Word document."}),
        filename="fake.docx",
        media_type=DOCX_MEDIA_TYPE,
    )

    assert response.status_code == 415
    assert response.json() == {"detail": "Document type does not match its content"}


def test_document_extraction_rejects_archive_expansion_bomb() -> None:
    entries = read_zip_entries(
        make_docx("Ordinary document text with a malicious compressed member.")
    )
    entries["word/media/padding.txt"] = b"0" * (2 * 1024 * 1024)

    response = post_document(
        make_zip(entries),
        filename="expansion.docx",
        media_type=DOCX_MEDIA_TYPE,
    )

    assert response.status_code == 413
    assert response.json() == {"detail": "Document archive exceeds safe expansion limits"}


def test_document_extraction_rejects_archive_path_traversal() -> None:
    entries = read_zip_entries(make_docx("Ordinary document text."))
    entries["../outside.xml"] = b"<outside />"

    response = post_document(
        make_zip(entries),
        filename="traversal.docx",
        media_type=DOCX_MEDIA_TYPE,
    )

    assert response.status_code == 422
    assert response.json() == {"detail": "Document package is unsafe"}


def test_document_extraction_rejects_macro_payloads() -> None:
    entries = read_zip_entries(make_docx("Ordinary document text."))
    entries["word/vbaProject.bin"] = b"macro payload"

    response = post_document(
        make_zip(entries),
        filename="macro.docx",
        media_type=DOCX_MEDIA_TYPE,
    )

    assert response.status_code == 422
    assert response.json() == {"detail": "Document package is unsafe"}


def test_document_extraction_accepts_exact_character_limit() -> None:
    large_text = "".join(f"{index:08x}" for index in range(31_250))
    content = make_docx(large_text)

    response = post_document(
        content,
        filename="maximum-text.docx",
        media_type=DOCX_MEDIA_TYPE,
        kind="jd",
    )

    assert response.status_code == 200
    assert response.json()["characterCount"] == 250_000


def test_document_extraction_enforces_character_limit() -> None:
    large_text = "".join(f"{index:08x}" for index in range(31_251))[:250_001]
    content = make_docx(large_text)

    response = post_document(
        content,
        filename="oversized-text.docx",
        media_type=DOCX_MEDIA_TYPE,
        kind="jd",
    )

    assert response.status_code == 413
    assert response.json() == {"detail": "Extracted text exceeds the 250000 character limit"}
