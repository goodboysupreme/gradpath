from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from pydantic import ValidationError

from app.schemas.documents import (
    DocumentExtractionResponse,
    DocumentFormat,
    DocumentKind,
)
from app.services.documents import (
    DocumentLimitError,
    DocumentParseError,
    DocumentTypeMismatch,
    EncryptedDocumentError,
    NoNativeTextError,
    UnsafeDocumentPackageError,
    extract_document,
    resolve_document_format,
)
from tests.test_documents import make_docx, make_pdf, make_zip, read_zip_entries


def test_docx_headers_body_and_footers_are_extracted_once_in_order() -> None:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Header contact information"
    document.add_paragraph("Main resume experience section")
    document.sections[0].footer.paragraphs[0].text = "Footer verification text"
    buffer = BytesIO()
    document.save(buffer)

    result = extract_document(
        buffer.getvalue(),
        DocumentFormat.DOCX,
        DocumentKind.RESUME,
    )

    assert result.text.count("Header contact information") == 1
    assert result.text.count("Main resume experience section") == 1
    assert result.text.count("Footer verification text") == 1
    assert result.text.index("Header contact information") < result.text.index(
        "Main resume experience section"
    )
    assert result.text.index("Main resume experience section") < result.text.index(
        "Footer verification text"
    )


def test_pure_pdf_extraction_and_limits() -> None:
    result = extract_document(
        make_pdf(["Native FastAPI resume evidence with enough text for extraction."]),
        DocumentFormat.PDF,
        DocumentKind.RESUME,
    )

    assert result.page_count == 1
    assert "FastAPI" in result.text

    with pytest.raises(DocumentLimitError, match="resume page limit"):
        extract_document(
            make_pdf([f"Native resume page {index} text." for index in range(11)]),
            DocumentFormat.PDF,
            DocumentKind.RESUME,
        )
    with pytest.raises(NoNativeTextError, match="OCR is not enabled"):
        extract_document(make_pdf([""]), DocumentFormat.PDF, DocumentKind.RESUME)
    with pytest.raises(EncryptedDocumentError, match="Encrypted documents"):
        extract_document(
            make_pdf(["Private encrypted resume text."], password="secret"),
            DocumentFormat.PDF,
            DocumentKind.RESUME,
        )
    with pytest.raises(DocumentParseError, match="could not be parsed"):
        extract_document(
            b"%PDF-1.7\ncorrupt",
            DocumentFormat.PDF,
            DocumentKind.RESUME,
        )


def test_pure_docx_extraction_includes_tables_and_rejects_unsafe_packages() -> None:
    content = make_docx(
        "Software internship description",
        table_cells=("Required skill", "PostgreSQL"),
    )
    result = extract_document(content, DocumentFormat.DOCX, DocumentKind.JD)

    assert result.page_count is None
    assert "Required skill\tPostgreSQL" in result.text

    traversal_entries = read_zip_entries(content)
    traversal_entries["../outside.xml"] = b"<outside />"
    with pytest.raises(UnsafeDocumentPackageError, match="package is unsafe"):
        extract_document(
            make_zip(traversal_entries),
            DocumentFormat.DOCX,
            DocumentKind.JD,
        )

    expansion_entries = read_zip_entries(content)
    expansion_entries["word/media/padding.txt"] = b"0" * (2 * 1024 * 1024)
    with pytest.raises(DocumentLimitError, match="safe expansion limits"):
        extract_document(
            make_zip(expansion_entries),
            DocumentFormat.DOCX,
            DocumentKind.JD,
        )

    with pytest.raises(DocumentParseError, match="could not be parsed"):
        extract_document(
            b"PK\x03\x04truncated",
            DocumentFormat.DOCX,
            DocumentKind.JD,
        )


def test_document_format_resolution_rejects_mismatches() -> None:
    docx = make_docx("Valid Word package content.")

    assert (
        resolve_document_format("resume.docx", "application/octet-stream", docx)
        is DocumentFormat.DOCX
    )
    with pytest.raises(DocumentTypeMismatch):
        resolve_document_format("resume.pdf", "application/pdf", docx)
    with pytest.raises(DocumentTypeMismatch):
        resolve_document_format("resume.pdf", "application/pdf", b"not a document")


def test_docx_main_content_type_and_response_metadata_are_enforced() -> None:
    content = make_docx("Ordinary Word document content.")
    with ZipFile(BytesIO(content)) as archive:
        entries = {info.filename: archive.read(info) for info in archive.infolist()}
    entries["[Content_Types].xml"] = entries["[Content_Types].xml"].replace(
        b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
        b"application/vnd.ms-word.document.macroEnabled.main+xml",
    )

    with pytest.raises(DocumentTypeMismatch):
        extract_document(
            make_zip(entries),
            DocumentFormat.DOCX,
            DocumentKind.JD,
        )

    with pytest.raises(ValidationError):
        DocumentExtractionResponse(
            document_kind=DocumentKind.RESUME,
            document_format=DocumentFormat.PDF,
            text="Native text is long enough.",
            page_count=None,
            character_count=27,
        )
    with pytest.raises(ValidationError):
        DocumentExtractionResponse(
            document_kind=DocumentKind.JD,
            document_format=DocumentFormat.DOCX,
            text="Native text is long enough.",
            page_count=1,
            character_count=27,
        )
